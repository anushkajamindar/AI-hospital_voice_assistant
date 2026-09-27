"""
Frontend API Blueprint
Provides REST endpoints for the MediCall dashboard frontend.
Mount this in your Flask app: app.register_blueprint(frontend_bp)

Routes:
  GET  /                         → serve dashboard HTML
  GET  /api/appointments         → list all appointments from Excel
  POST /api/appointments         → add a new appointment manually
  DELETE /api/appointments/<idx> → delete appointment by row index (0-based)
  GET  /api/export               → download the Excel file
"""

import os
from flask import Blueprint, jsonify, request, send_file, send_from_directory
from datetime import datetime

frontend_bp = Blueprint("frontend", __name__)

# ── CORS helper – allows the HTML to call the API even when opened as file:// ──
def _corsify(response):
    response.headers["Access-Control-Allow-Origin"] = "*"
    response.headers["Access-Control-Allow-Methods"] = "GET, POST, DELETE, OPTIONS"
    response.headers["Access-Control-Allow-Headers"] = "Content-Type"
    return response

@frontend_bp.after_request
def after_request(response):
    return _corsify(response)

# ── Resolve paths ──────────────────────────────────────────────────────────
_HERE = os.path.dirname(os.path.abspath(__file__))

def _excel_path():
    try:
        from app.services.config import Config
        return Config.EXCEL_FILE_PATH
    except Exception:
        return os.path.join(_HERE, "appointments.xlsx")


# ── Serve the dashboard ────────────────────────────────────────────────────
@frontend_bp.route("/", methods=["GET"])
def dashboard():
    """Serve the main frontend HTML."""
    return send_from_directory(_HERE, "app.html")


# ── GET /api/appointments ─────────────────────────────────────────────────
@frontend_bp.route("/api/appointments", methods=["GET"])
def get_appointments():
    """Return all appointments from the Excel file as JSON."""
    path = _excel_path()
    appointments = []

    if not os.path.exists(path):
        return jsonify({"appointments": [], "total": 0})

    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Appointments"]

        # Read header row to determine column order
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

        # Build column index map (case-insensitive)
        col_map = {str(h).strip().lower(): i for i, h in enumerate(headers) if h}

        def _col(name):
            """Get cell value by column name."""
            idx = col_map.get(name.lower())
            return idx

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue

            # Map columns by header name for robustness
            def get(name, fallback=None):
                idx = col_map.get(name.lower())
                if idx is None:
                    return fallback
                return row[idx] if idx < len(row) else fallback

            # Support both old (swapped) and new (correct) column order
            serial         = get("#") or get("serial")
            name           = get("patient name")
            disease        = get("disease / symptoms") or get("disease/symptoms") or get("disease")
            doctor         = get("doctor preference") or get("doctor name") or get("doctor")
            preferred_time = get("preferred time")
            call_sid       = get("call sid")
            recorded_at    = get("recorded at")

            # Convert datetime objects to strings
            if isinstance(recorded_at, datetime):
                recorded_at = recorded_at.strftime("%Y-%m-%d %H:%M:%S")
            if isinstance(preferred_time, datetime):
                preferred_time = preferred_time.strftime("%Y-%m-%d %H:%M")

            appointments.append({
                "serial":         int(serial) if serial else len(appointments) + 1,
                "name":           str(name)   if name   else "Unknown",
                "disease":        str(disease) if disease else "",
                "doctor":         str(doctor)  if doctor  else "",
                "preferred_time": str(preferred_time) if preferred_time else "",
                "call_sid":       str(call_sid) if call_sid else "",
                "recorded_at":    str(recorded_at) if recorded_at else "",
            })

        wb.close()

    except Exception as e:
        return jsonify({"error": str(e), "appointments": [], "total": 0}), 500

    return jsonify({"appointments": appointments, "total": len(appointments)})


# ── POST /api/appointments ────────────────────────────────────────────────
@frontend_bp.route("/api/appointments", methods=["POST"])
def add_appointment():
    """Manually add a new appointment (from frontend form)."""
    data = request.get_json(force=True)

    name           = str(data.get("name", "")).strip()
    disease        = str(data.get("disease", "")).strip()
    doctor         = str(data.get("doctor", "")).strip()
    preferred_time = str(data.get("preferred_time", "")).strip()
    call_sid       = str(data.get("call_sid", f"MANUAL-{int(datetime.now().timestamp())}")).strip()

    if not name or not disease or not doctor or not preferred_time:
        return jsonify({"error": "Missing required fields: name, disease, doctor, preferred_time"}), 400

    try:
        from app.services.excel_service import save_appointment
        save_appointment(
            name=name,
            disease=disease,
            doctor=doctor,
            preferred_time=preferred_time,
            call_sid=call_sid,
        )
    except Exception as e:
        return jsonify({"error": f"Failed to save: {e}"}), 500

    return jsonify({"status": "ok", "message": f"Appointment saved for {name}"}), 201


# ── DELETE /api/appointments/<idx> ───────────────────────────────────────
@frontend_bp.route("/api/appointments/<int:idx>", methods=["DELETE"])
def delete_appointment(idx):
    """
    Delete appointment at the given 0-based index (row = idx + 2 in Excel).
    """
    path = _excel_path()
    if not os.path.exists(path):
        return jsonify({"error": "Excel file not found"}), 404

    try:
        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb["Appointments"]

        # Row 1 = header, row 2 = first data row (idx=0)
        excel_row = idx + 2
        if excel_row < 2 or excel_row > ws.max_row:
            return jsonify({"error": f"Row index {idx} out of range"}), 400

        ws.delete_rows(excel_row)

        # Re-sequence the serial numbers
        for i, row in enumerate(ws.iter_rows(min_row=2), start=1):
            if row[0].value is not None:
                row[0].value = i

        wb.save(path)
        wb.close()

    except Exception as e:
        return jsonify({"error": str(e)}), 500

    return jsonify({"status": "ok", "message": f"Appointment {idx} deleted"})


# ── GET /api/export ───────────────────────────────────────────────────────
@frontend_bp.route("/api/export", methods=["GET"])
def export_excel():
    """Download the raw Excel file."""
    path = _excel_path()
    if not os.path.exists(path):
        return jsonify({"error": "No Excel file found"}), 404

    return send_file(
        path,
        as_attachment=True,
        download_name="appointments.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
